import sqlite3
import pandas as pd
from pathlib import Path
from openpyxl.utils import get_column_letter
from data.tipos import tipos_eq, tipos_item, tipos_setor, tipos_status, tipos_status_calibr
from helpers import DB_FILE, get_connection

def exportar_para_excel(file_path="data/excel_output/equipamentos_itens.xlsx"):
    conn = get_connection(DB_FILE)
    
    # ----------------------------- Query Equipamentos -----------------------------
    query_equip = "SELECT * FROM equipamentos"
    df_equip = pd.read_sql_query(query_equip, conn)

    # ----------------------------- Adicionar coluna "Valor" -----------------------------
    query_valor = """
        SELECT equipamento_id, valor
        FROM ciclo_vida
        WHERE tipo_item_id = 9
    """
    df_valor = pd.read_sql_query(query_valor, conn)
    df_valor = df_valor.groupby("equipamento_id", as_index=False).first()  # caso tenha mais de um registro
    df_equip = df_equip.merge(df_valor, left_on="id", right_on="equipamento_id", how="left")
    df_equip.drop(columns=["equipamento_id"], inplace=True)
    df_equip.rename(columns={"valor": "Valor"}, inplace=True)

    # ----------------------------- Mapeamentos -----------------------------
    status_dict = dict(tipos_status)
    status_calibr_dict = dict(tipos_status_calibr)
    setor_dict = dict(tipos_setor)
    tipo_eq_dict = {i+1: nome for i, nome in enumerate(tipos_eq)}

    df_equip['tipo_eq_id'] = df_equip['tipo_eq_id'].map(tipo_eq_dict)
    df_equip['status_id'] = df_equip['status_id'].map(status_dict)
    df_equip['status_calibracao_id'] = df_equip['status_calibracao_id'].map(status_calibr_dict)
    df_equip['setor_id'] = df_equip['setor_id'].map(setor_dict)

    # Datas (somente dia, mês e ano)
    df_equip['data_aquisicao'] = pd.to_datetime(df_equip['data_aquisicao'], errors='coerce').dt.strftime('%d-%m-%Y')
    df_equip['ultima_calibracao'] = pd.to_datetime(df_equip['ultima_calibracao'], errors='coerce').dt.strftime('%d-%m-%Y')

    # Renomear colunas
    df_equip.rename(columns={
        'nome_eq': 'Nome',
        'tipo_eq_id': 'Tipo',
        'sigla_eq': 'Sigla',
        'setor_id': 'Setor',
        'status_id': 'Status',
        'sond_id': 'Sond',
        'data_aquisicao': 'Data Aquisição',
        'ultima_calibracao': 'Última Calibração',
        'periodicidade': 'Periodicidade',
        'status_calibracao_id': 'Status Calibração',
        'fabricante': 'Fabricante',
        'modelo': 'Modelo',
        'modelo_tecnico': 'Modelo Técnico',
        'numero_serie': 'Número de Série',
        'extra_info': 'Informações Extras'
    }, inplace=True)

    # ----------------------------- Query Ciclo de Vida -----------------------------
    query_itens = """
    SELECT cv.id, cv.equipamento_id, cv.tipo_item_id, cv.descricao, cv.info_especial,
           cv.data, cv.fornecedor, cv.valor, e.nome_eq
    FROM ciclo_vida cv
    JOIN equipamentos e ON cv.equipamento_id = e.id
    """
    df_itens = pd.read_sql_query(query_itens, conn)

    tipos_item_dict = dict(tipos_item)
    df_itens['tipo_item_id'] = df_itens['tipo_item_id'].map(tipos_item_dict)

    # Datas com hora e minuto
    df_itens['data'] = pd.to_datetime(df_itens['data'], errors='coerce').dt.strftime('%d-%m-%Y %H:%M')

    df_itens.rename(columns={
        'id': 'ID',
        'equipamento_id': 'Equipamento ID',
        'nome_eq': 'Equipamento',
        'tipo_item_id': 'Tipo Item',
        'descricao': 'Descrição',
        'info_especial': 'Info Especial',
        'data': 'Data',
        'fornecedor': 'Fornecedor',
        'valor': 'Valor'
    }, inplace=True)

    # Reordenar colunas
    cols = df_itens.columns.tolist()
    if 'Equipamento' in cols:
        cols.remove('Equipamento')
        cols = cols[:1] + ['Equipamento'] + cols[1:]
    df_itens = df_itens[cols]

    # ----------------------------- Exportar para Excel -----------------------------
    file_path = Path(file_path)

    # Reordenar colunas: colocar "Valor" antes de "Informações Extras"
    cols = df_equip.columns.tolist()
    if "Valor" in cols and "Informações Extras" in cols:
        cols.insert(cols.index("Informações Extras"), cols.pop(cols.index("Valor")))
        df_equip = df_equip[cols]


    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        df_equip.to_excel(writer, sheet_name='Equipamentos', index=False)
        df_itens.to_excel(writer, sheet_name='Itens Ciclo de Vida', index=False)
        
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]
            worksheet.auto_filter.ref = worksheet.dimensions

            # Ajustar largura automática
            for col_idx, col in enumerate(worksheet.iter_cols(1, worksheet.max_column), start=1):
                max_length = 0
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                worksheet.column_dimensions[get_column_letter(col_idx)].width = max_length + 2

    print(f"Arquivo Excel exportado com sucesso: {file_path}")


# ========================== NOVAS FUNÇÕES ==========================

def exportar_eq_csv(file_path="data/csv_output/equipamentos.csv"):
    """Exporta os equipamentos para um arquivo CSV"""
    conn = get_connection(DB_FILE)
    query_equip = "SELECT * FROM equipamentos"
    df_equip = pd.read_sql_query(query_equip, conn)

    # ----------------------------- Adicionar coluna "Valor" -----------------------------
    query_valor = """
        SELECT equipamento_id, valor
        FROM ciclo_vida
        WHERE tipo_item_id = 9
    """
    df_valor = pd.read_sql_query(query_valor, conn)
    df_valor = df_valor.groupby("equipamento_id", as_index=False).first()
    df_equip = df_equip.merge(df_valor, left_on="id", right_on="equipamento_id", how="left")
    df_equip.drop(columns=["equipamento_id"], inplace=True)
    df_equip.rename(columns={"valor": "Valor"}, inplace=True)

    # ----------------------------- Mapear IDs para nomes -----------------------------
    df_equip['tipo_eq_id'] = df_equip['tipo_eq_id'].map({i+1: nome for i, nome in enumerate(tipos_eq)})
    df_equip['status_id'] = df_equip['status_id'].map(dict(tipos_status))
    df_equip['status_calibracao_id'] = df_equip['status_calibracao_id'].map(dict(tipos_status_calibr))
    df_equip['setor_id'] = df_equip['setor_id'].map(dict(tipos_setor))

    # Datas (somente dia, mês e ano)
    df_equip['data_aquisicao'] = pd.to_datetime(df_equip['data_aquisicao'], errors='coerce').dt.strftime('%d-%m-%Y')
    df_equip['ultima_calibracao'] = pd.to_datetime(df_equip['ultima_calibracao'], errors='coerce').dt.strftime('%d-%m-%Y')

    # Renomear colunas
    df_equip.rename(columns={
        'nome_eq': 'Nome',
        'tipo_eq_id': 'Tipo',
        'sigla_eq': 'Sigla',
        'setor_id': 'Setor',
        'status_id': 'Status',
        'sond_id': 'Sond',
        'data_aquisicao': 'Data Aquisição',
        'ultima_calibracao': 'Última Calibração',
        'periodicidade': 'Periodicidade',
        'status_calibracao_id': 'Status Calibração',
        'fabricante': 'Fabricante',
        'modelo': 'Modelo',
        'modelo_tecnico': 'Modelo Técnico',
        'numero_serie': 'Número de Série',
        'extra_info': 'Informações Extras'
    }, inplace=True)

    # ----------------------------- Exportar CSV -----------------------------
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    df_equip.to_csv(file_path, index=False, encoding="utf-8-sig", sep=";")
    print(f"Arquivo CSV de equipamentos exportado com sucesso: {file_path}")


def exportar_item_csv(file_path="data/csv_output/itens_ciclo_vida.csv"):
    """Exporta os itens do ciclo de vida para um arquivo CSV"""
    conn = get_connection(DB_FILE)
    query_itens = """
    SELECT cv.id, cv.equipamento_id, cv.tipo_item_id, cv.descricao, cv.info_especial,
           cv.data, cv.fornecedor, cv.valor, e.nome_eq
    FROM ciclo_vida cv
    JOIN equipamentos e ON cv.equipamento_id = e.id
    """
    df_itens = pd.read_sql_query(query_itens, conn)

    # Mapear tipo item
    df_itens['tipo_item_id'] = df_itens['tipo_item_id'].map(dict(tipos_item))

    # Converter datas com hora e minuto
    df_itens['data'] = pd.to_datetime(df_itens['data'], errors='coerce').dt.strftime('%d-%m-%Y %H:%M')

    # Renomear colunas
    df_itens.rename(columns={
        'id': 'ID',
        'equipamento_id': 'Equipamento ID',
        'nome_eq': 'Equipamento',
        'tipo_item_id': 'Tipo Item',
        'descricao': 'Descrição',
        'info_especial': 'Info Especial',
        'data': 'Data',
        'fornecedor': 'Fornecedor',
        'valor': 'Valor'
    }, inplace=True)

    # Reordenar colunas
    cols = df_itens.columns.tolist()
    if 'Equipamento' in cols:
        cols.remove('Equipamento')
        cols = cols[:1] + ['Equipamento'] + cols[1:]
    df_itens = df_itens[cols]

    # ----------------------------- Exportar CSV -----------------------------
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    df_itens.to_csv(file_path, index=False, encoding="utf-8-sig", sep=";")
    print(f"Arquivo CSV de itens exportado com sucesso: {file_path}")
