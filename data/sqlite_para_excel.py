import sqlite3
import pandas as pd
from pathlib import Path
from openpyxl.utils import get_column_letter
from data.tipos import tipos_eq, tipos_item, tipos_setor, tipos_status, tipos_status_calibr
from helpers import DB_FILE, get_connection

def exportar_para_excel(file_path="data/excel_output/equipamentos_itens.xlsx"):
    conn = get_connection(DB_FILE)
    
    # ----------------------------- Query Equipamentos -----------------------------
    query_equip = "SELECT * FROM equipamentos"
    df_equip = pd.read_sql_query(query_equip, conn)
    
    status_dict = dict(tipos_status)
    status_calibr_dict = dict(tipos_status_calibr)
    setor_dict = dict(tipos_setor)
    tipo_eq_dict = {i+1: nome for i, nome in enumerate(tipos_eq)}
    
    df_equip['tipo_eq_id'] = df_equip['tipo_eq_id'].map(tipo_eq_dict)
    df_equip['status_id'] = df_equip['status_id'].map(status_dict)
    df_equip['status_calibracao_id'] = df_equip['status_calibracao_id'].map(status_calibr_dict)
    df_equip['setor_id'] = df_equip['setor_id'].map(setor_dict)
    
    df_equip['data_aquisicao'] = pd.to_datetime(df_equip['data_aquisicao'], errors='coerce').dt.strftime('%d-%m-%Y')
    df_equip['ultima_calibracao'] = pd.to_datetime(df_equip['ultima_calibracao'], errors='coerce').dt.strftime('%d-%m-%Y')
    
    df_equip.rename(columns={
        'nome_eq': 'Nome',
        'tipo_eq_id': 'Tipo',
        'sigla_eq': 'Sigla',
        'setor_id': 'Setor',
        'status_id': 'Status',
        'sond_id': 'Sond',
        'data_aquisicao': 'Data Aquisição',
        'ultima_calibracao': 'Última Calibração',
        'periodicidade': 'Periodicidade',
        'status_calibracao_id': 'Status Calibração',
        'fabricante': 'Fabricante',
        'modelo': 'Modelo',
        'modelo_tecnico': 'Modelo Técnico',
        'numero_serie': 'Número de Série',
        'extra_info': 'Informações Extras'
    }, inplace=True)
    
    # ----------------------------- Query Ciclo de Vida -----------------------------
    query_itens = """
    SELECT cv.id, cv.equipamento_id, cv.tipo_item_id, cv.descricao, cv.info_especial,
           cv.data, cv.fornecedor, cv.valor, e.nome_eq
    FROM ciclo_vida cv
    JOIN equipamentos e ON cv.equipamento_id = e.id
    """
    df_itens = pd.read_sql_query(query_itens, conn)
    
    tipos_item_dict = dict(tipos_item)
    df_itens['tipo_item_id'] = df_itens['tipo_item_id'].map(tipos_item_dict)
    
    df_itens['data'] = pd.to_datetime(df_itens['data'], errors='coerce').dt.strftime('%d-%m-%Y')
    
    df_itens.rename(columns={
        'id': 'ID',
        'equipamento_id': 'Equipamento ID',
        'nome_eq': 'Equipamento',
        'tipo_item_id': 'Tipo Item',
        'descricao': 'Descrição',
        'info_especial': 'Info Especial',
        'data': 'Data',
        'fornecedor': 'Fornecedor',
        'valor': 'Valor'
    }, inplace=True)
    
    # Reordenar colunas: segunda coluna será "Equipamento"
    cols = df_itens.columns.tolist()
    if 'Equipamento' in cols:
        cols.remove('Equipamento')
        cols = cols[:1] + ['Equipamento'] + cols[1:]
    df_itens = df_itens[cols]
    
    # ----------------------------- Exportar para Excel -----------------------------
    file_path = Path(file_path)
    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        df_equip.to_excel(writer, sheet_name='Equipamentos', index=False)
        df_itens.to_excel(writer, sheet_name='Itens Ciclo de Vida', index=False)
        
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]
            worksheet.auto_filter.ref = worksheet.dimensions
    
    print(f"Arquivo exportado com sucesso: {file_path}")

    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        df_equip.to_excel(writer, sheet_name='Equipamentos', index=False)
        df_itens.to_excel(writer, sheet_name='Itens Ciclo de Vida', index=False)
        
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]
            worksheet.auto_filter.ref = worksheet.dimensions
            
            # Ajustar largura das colunas pelo tamanho do header
            for col_idx, col in enumerate(worksheet.iter_cols(1, worksheet.max_column), start=1):
                max_length = 0
                header = col[0].value
                if header:
                    max_length = len(str(header))
                column_letter = get_column_letter(col_idx)
                worksheet.column_dimensions[column_letter].width = max_length + 2  # +2 para espaçamento extra

    print(f"Arquivo formatado com sucesso")

